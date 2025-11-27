from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from uuid import uuid4
import io
import time
from app.database import performance_collection
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from app.models import SentimentRequest, SentimentResponse
from app.sentiment_service import analyze_text
from datetime import datetime
from app.database import collection, users_collection, activity_collection, files_collection
from app.auth import oauth2_scheme, jwt, SECRET_KEY, ALGORITHM
from app.blob_service import (
    upload_bytes,
    delete_blob,
    download_bytes,
    list_user_blobs,
    generate_report_sas,
)
from app.email_service import send_azure_email


router = APIRouter()

def verify_token(token: str = Depends(oauth2_scheme)):
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username = payload.get("sub")
        if not username:
            raise HTTPException(status_code=401, detail="Invalid token")
        return username
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid token")


@router.post("/upload", tags=["File Uploads"], summary="Upload CSV to Blob Storage")
async def upload_file(
    file: UploadFile = File(...),
    username: str = Depends(verify_token),
):
    """Upload a CSV file. Returns a generated file_id."""
    if not file.filename.lower().endswith(".csv"):
        raise HTTPException(status_code=400, detail="Only CSV files allowed")

    content = await file.read()

    file_id = str(uuid4())
    blob_path = f"{username}/uploads/{file_id}.csv"

    upload_bytes(content, blob_path)

    # 🔹 Log this upload in MongoDB
    activity_collection.insert_one({
        "username": username,
        "event": "file_uploaded",
        "file_id": file_id,
        "filename": file.filename,
        "blob_path": blob_path,
        "size_bytes": len(content),
        "timestamp": datetime.utcnow(),
    })
    user_doc = users_collection.find_one({"username": username})
    email = user_doc["email"] if user_doc else None

    files_collection.insert_one({
    "file_id": file_id,
    "username": username,
    "email": email,
    "filename": file.filename,
    "blob_path": blob_path,
    "size_bytes": len(content),
    "uploaded_at": datetime.utcnow()
     })

    return {"message": "Upload successful", "file_id": file_id}


@router.get("/file", tags=["File Uploads"], summary="Get uploaded file info")
def get_file_info(file_id: str, username: str = Depends(verify_token)):
    """Check if an uploaded file exists for this user."""
    blob_path = f"{username}/uploads/{file_id}.csv"

    try:
        download_bytes(blob_path)
    except Exception:
        raise HTTPException(status_code=404, detail="File not found")

    return {"file_id": file_id, "status": "Available"}


@router.delete("/file", tags=["File Uploads"], summary="Delete uploaded CSV")
def delete_uploaded_file(file_id: str, username: str = Depends(verify_token)):
    """Delete an uploaded CSV (not the summary) using file_id."""
    blob_path = f"{username}/uploads/{file_id}.csv"

    try:
        delete_blob(blob_path)
    except Exception:
        raise HTTPException(status_code=404, detail="File not found")
    
    activity_collection.insert_one({
        "username": username,
        "event": "file_deleted",
        "file_id": file_id,
        "blob_path": blob_path,
        "timestamp": datetime.utcnow(),
    })

    return {"message": " Your .CSV file deleted successful", "file_id": file_id}


@router.get(
    "/list",
    tags=["File Uploads"],
    summary="List all uploaded CSV and summary files (by file_id)",
)
def list_files(username: str = Depends(verify_token)):
    """
    List CSV uploads and Excel summaries.

    Returns just IDs:
      uploads   -> list of file_ids that exist under /uploads/
      summaries -> list of file_ids that have a summary under /results/
    """
    upload_blobs = list_user_blobs(f"{username}/uploads/")
    summary_blobs = list_user_blobs(f"{username}/results/")

    uploads = [b.split("/")[-1].replace(".csv", "") for b in upload_blobs]
    summaries = [b.split("/")[-1].replace("_summary.xlsx", "") for b in summary_blobs]

    return {"uploads": uploads, "summaries": summaries}

@router.post("/analyze", tags=["Analyze"], response_model=SentimentResponse)
def analyze_sentiment(request: SentimentRequest, username: str = Depends(verify_token)):

    existing = collection.find_one({
        "username": username,
        "text": request.text
    })
    if existing:
        raise HTTPException(
            status_code=409, detail="Duplicate: This text is already analyzed"
        )

    result = analyze_text(request.text)

    user_doc = users_collection.find_one({"username": username})
    email = user_doc["email"] if user_doc else None

    collection.insert_one({
        "username": username,
        "email": email,
        "text": request.text,
        "label": result["label"],
        "score": float(result["score"]),
        "timestamp": datetime.utcnow()
    })

    return {
        "text": request.text,
        "label": result["label"],
        "score": result["score"]
    }
  

@router.get("/history", tags=["Analyze"])
def history(username: str = Depends(verify_token)):
    """Get analysis history for the logged-in user."""
    docs = list(collection.find({"username": username}, {"_id": 0, "username": 0, "email":0, "timestamp":0}))
    return {"history": docs}


@router.delete("/delete", tags=["Analyze"])
def delete_text(text: str, username: str = Depends(verify_token)):
    """Delete a single text analysis entry."""
    result = collection.delete_one({"username": username, "text": text})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Text not found")
    return {"username": username, "text": text}


@router.post(
    "/analyze",
    tags=["Analyze"],
    summary="Analyze an uploaded CSV by file_id and generate Excel report",
)
def analyze_uploaded_file(file_id: str, username: str = Depends(verify_token)):
    """
    Analyze a previously uploaded CSV (by file_id), generate a modern Excel report
    with two sheets, upload it to Blob, and send a plain-text email with a SAS link.
    """
    start_time = time.time()

    csv_blob = f"{username}/uploads/{file_id}.csv"

    try:
        file_bytes = download_bytes(csv_blob)
    except Exception:
        raise HTTPException(status_code=404, detail="CSV not found")

    df = pd.read_csv(io.BytesIO(file_bytes))
    if "text" not in df.columns:
        raise HTTPException(status_code=400, detail="CSV must contain a 'text' column")

    texts = df["text"].dropna().tolist()

    results = []
    for t in texts:
        res = analyze_text(t)
        results.append(
            {"text": t, "label": res["label"], "score": float(res["score"])}
        )

    total = len(results)
    positive = sum(r["label"] == "POSITIVE" for r in results)
    negative = sum(r["label"] == "NEGATIVE" for r in results)
    neutral = sum(r["label"] == "NEUTRAL" for r in results)

    if total == 0:
        pos_pct = neg_pct = neu_pct = 0.0
    else:
        pos_pct = round(positive / total * 100, 2)
        neg_pct = round(negative / total * 100, 2)
        neu_pct = round(neutral / total * 100, 2)


    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Report Summary"

    ws_summary["A1"] = "Sentiment Analysis Report"
    ws_summary["A3"] = "File ID"
    ws_summary["B3"] = file_id
    ws_summary["A4"] = "Total Rows"
    ws_summary["B4"] = total

    ws_summary["A6"] = "Sentiment"
    ws_summary["B6"] = "Count"
    ws_summary["C6"] = "Percentage (%)"

    ws_summary["A7"] = "POSITIVE"
    ws_summary["B7"] = positive
    ws_summary["C7"] = pos_pct

    ws_summary["A8"] = "NEGATIVE"
    ws_summary["B8"] = negative
    ws_summary["C8"] = neg_pct

    ws_summary["A9"] = "NEUTRAL"
    ws_summary["B9"] = neutral
    ws_summary["C9"] = neu_pct

    chart = BarChart()
    chart.title = "Sentiment Percentage Distribution"
    chart.x_axis.title = "Sentiment"
    chart.y_axis.title = "Percentage (%)"

    data = Reference(ws_summary, min_col=3, min_row=6, max_row=9)  # C6:C9
    categories = Reference(ws_summary, min_col=1, min_row=7, max_row=9)  # A7:A9

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    ws_summary.add_chart(chart, "E4")

    ws_data = wb.create_sheet(title="Raw Data")
    ws_data.append(["text", "label", "score"])
    for r in results:
        ws_data.append([r["text"], r["label"], r["score"]])

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_bytes = excel_stream.getvalue()

    summary_blob = f"{username}/results/{file_id}_summary.xlsx"
    upload_bytes(excel_bytes, summary_blob)

    user_doc = users_collection.find_one({"username": username})
    if user_doc:
        sas_link = generate_report_sas(summary_blob, expiry_minutes=60)

        email_body = f"""Hello {username},

Your sentiment analysis report is ready.

File ID: {file_id}
Total Rows: {total}

Sentiment Distribution:
POSITIVE: {pos_pct}%
NEGATIVE: {neg_pct}%
NEUTRAL: {neu_pct}%

Download your report (valid 60 minutes):
{sas_link}

Regards,
Cloud Sentiment Analysis Platform
"""

        send_azure_email(
            to_email=user_doc["email"],
            subject="Your Sentiment Analysis Report is Ready",
            body=email_body,
        )

    activity_collection.insert_one({
        "username": username,
        "event": "file_analyzed",
        "file_id": file_id,
        "csv_blob": csv_blob,
        "summary_blob": summary_blob,
        "total_rows": total,
        "positive_count": positive,
        "negative_count": negative,
        "neutral_count": neutral,
        "positive_pct": pos_pct,
        "negative_pct": neg_pct,
        "neutral_pct": neu_pct,
        "timestamp": datetime.utcnow(),
    })

    latency_ms = round((time.time() - start_time) * 1000, 3)

    performance_collection.insert_one({
        "type": "file_analysis_latency",
        "username": username,
        "file_id": file_id,
        "total_rows": total,
        "latency_ms": latency_ms,
        "timestamp": datetime.utcnow()
    })

    return {
        "message": "Analysis completed and summary Excel generated",
        "file_id": file_id,
        "summary_available": True,
    }

@router.get(
    "/download",
    tags=["Analyze"],
    summary="Download the summary Excel by file_id",
)
def download_summary(file_id: str, username: str = Depends(verify_token)):
    """Download the Excel summary using only file_id."""
    blob = f"{username}/results/{file_id}_summary.xlsx"

    try:
        file_bytes = download_bytes(blob)
    except Exception:
        raise HTTPException(status_code=404, detail="Summary file not found")

    return StreamingResponse(
        io.BytesIO(file_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{file_id}_summary.xlsx"'
        },
    )


@router.delete(
    "/deleteresult",
    tags=["Analyze"],
    summary="Delete the summary Excel by file_id",
)
def delete_summary(file_id: str, username: str = Depends(verify_token)):
    """Delete the generated summary Excel using only file_id."""
    blob = f"{username}/results/{file_id}_summary.xlsx"

    try:
        delete_blob(blob)
    except Exception:
        raise HTTPException(status_code=404, detail="Summary file not found")
    
    activity_collection.insert_one({
        "username": username,
        "event": "summary_deleted",
        "file_id": file_id,
        "summary_blob": blob,
        "timestamp": datetime.utcnow(),
    })

    return {"message": "Summary deleted", "file_id": file_id}
@router.get("/performance", tags=["Performance"])
def global_file_latency_violin(username: str = Depends(verify_token)):
    """
    Global violin plot showing latency distribution for ALL USERS.
    """
    import matplotlib.pyplot as plt
    import numpy as np


    docs = list(performance_collection.find({
        "type": "file_analysis_latency"
    }))  # <-- removed username filtering

    if not docs:
        raise HTTPException(status_code=404, detail="No global latency data found")

    latencies = [d["latency_ms"] for d in docs]

    plt.figure(figsize=(8, 5))
    plt.violinplot(latencies, showmeans=True, showmedians=True)
    plt.title("Global File Analysis Latency Distribution (ms)")
    plt.ylabel("Latency (ms)")

    buf = io.BytesIO()
    plt.savefig(buf, format="png")
    buf.seek(0)
    plt.close()

    return StreamingResponse(buf, media_type="image/png")


